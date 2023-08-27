import os
from typing import Generator

import latex2mathml.converter
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Twips
from docx.table import Table
from docx.text.run import Run
from lxml import etree

from ..layout_tracker import LayoutState
from ..renderable import Renderable
from ..rendered_info import RenderedInfo
from ..util import create_element


_HEIGHT = Pt(50)


class Formula(Renderable):
    def __init__(self, parent, latex_formula: str):
        try:
            mathml = latex2mathml.converter.convert(latex_formula)
            tree = etree.fromstring(mathml)
            xslt = etree.parse(
                os.path.join(os.path.dirname(__file__), "mml2omml")
            )
            transform = etree.XSLT(xslt)
            new_dom = transform(tree)
            word_math = new_dom.getroot()
        except Exception:
            raise ValueError(f"Can't parse the formula:\n{latex_formula}")

        sect = parent.part.document.sections[0]

        # todo: style inheritance
        left_margin = Twips(int(parent.part.styles["Normal Table"]._element.xpath("w:tblPr/w:tblCellMar/w:left")[0].attrib["{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"]))
        right_margin = Twips(int(parent.part.styles["Normal Table"]._element.xpath("w:tblPr/w:tblCellMar/w:right")[0].attrib["{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"]))

        table_width = sect.page_width - sect.right_margin - sect.left_margin + left_margin + right_margin

        self._table = table = Table(create_element("w:tbl", [
            create_element("w:tblGrid")
        ]), parent)

        row = table.add_row()
        right_cell_width = Pt(30)
        table.add_column(table_width - right_cell_width)
        table.add_column(right_cell_width)
        left_cell = table.cell(0,0)
        right_cell = table.cell(0,1)
        # TODO: implement proper size
        row.height = _HEIGHT

        left_paragraph = left_cell.paragraphs[0]
        left_paragraph.style = "Formula Content"
        left_paragraph._p.append(word_math)
        left_cell.vertical_alignment = \
            WD_CELL_VERTICAL_ALIGNMENT.CENTER

        right_paragraph = right_cell.paragraphs[0]
        right_paragraph.style = "Formula Numbering"
        right_paragraph._p.append(create_element("w:r", "("))
        self._numbering_run = create_element("w:r", "?")
        right_paragraph._p.append(
            create_element("w:fldSimple", {
                "w:instr": f"SEQ formula \\* ARABIC"
            }, [self._numbering_run]))
        right_paragraph._p.append(create_element("w:r", ")"))
        right_cell.vertical_alignment = \
            WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def render(self, previous_rendered: RenderedInfo, layout_state: LayoutState) -> Generator[
            "RenderedInfo | Renderable", None, None]:
        height = _HEIGHT

        if height > layout_state.remaining_page_height:
            height += layout_state.remaining_page_height

        yield RenderedInfo(self._table, height)
